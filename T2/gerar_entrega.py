"""T2 - gera os artefatos de entrega a partir de T2/resultados/*.json (não roda simulação).

Uso:
  python T2/gerar_entrega.py        # PPTX, XLSX, roteiro (MD e PDF) e VIDEO_PENDENTE.txt
  python T2/gerar_entrega.py zip    # T2/Entrega_T2.zip (exige entrega/Apresentacao_T2.pdf)
"""

import json
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
RESULTADOS = BASE / "resultados"
ENTREGA = BASE / "entrega"
CENARIOS = ("inicial", "melhorado")
FILAS = ("Q1", "Q2", "Q3")
PAPEIS = {"Q1": "Triagem", "Q2": "Reparo", "Q3": "Teste"}
INTEGRANTES = ("Artur Lemos Pereira", "Raul Neves Pedroso", "Humberto Corrêa Gomes")
MIT = "https://web.mit.edu/urban_or_book/www/book/chapter4/4.4.html"
IGNORAR = {".workflow", "temp", "tmp", "logs", "__pycache__"}

NAVY, TEAL, AMARELO = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x1A, 0x8C, 0x84), RGBColor(0xF2, 0xC2, 0x30)
CLARO, TEXTO, BRANCO = RGBColor(0xF5, 0xF7, 0xFA), RGBColor(0x1F, 0x29, 0x37), RGBColor(0xFF, 0xFF, 0xFF)
LISTRA = RGBColor(0xE8, 0xF1, 0xF4)

CHECKLIST = [
    "Ensaiar a apresentação inteira com cronômetro (alvo de 8 a 10 minutos).",
    "Abrir Apresentacao_T2.pptx no PowerPoint e iniciar a apresentação de slides.",
    "No Zoom, iniciar a reunião com os três integrantes e compartilhar a janela da apresentação.",
    "Conferir o áudio de cada integrante antes de gravar.",
    "Gravar a reunião; cada integrante fala nos slides indicados neste roteiro.",
    "Revisar a gravação inteira (áudio, slides visíveis, duração).",
    "Publicar no YouTube como não listado e testar o link em uma janela anônima.",
    "Criar T2/entrega/VIDEO.txt com o link real e apagar VIDEO_PENDENTE.txt.",
    "Reconstruir o ZIP (python T2/gerar_entrega.py zip) antes de enviar no Moodle.",
]


def br(x, d=2):
    """Número no formato pt-BR (milhar com ponto, decimal com vírgula)."""
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x, d=2):
    return br(100 * x, d) + "%"


def carregar(resultados=RESULTADOS):
    resultados = Path(resultados)
    return {c: json.loads((resultados / f"{c}.json").read_text(encoding="utf-8")) for c in CENARIOS}


def principal(dados_cenario):
    s = dados_cenario["semente_principal"]
    return next(r for r in dados_cenario["replicacoes"] if r["seed"] == s)


def _variacao(i, m):
    return br(100 * (m - i) / i, 2) + "%" if i else "–"


# ---------------------------------------------------------------- roteiro (falas)

def roteiro(dados):
    """Lista de slides: título, orador, segundos-alvo e fala (também vira nota do PPTX)."""
    ini, mel = dados["inicial"], dados["melhorado"]
    ri, rm = principal(ini), principal(mel)
    Ii, Im = ri["indices"], rm["indices"]
    q = ini["config"]["queues"]
    cm = mel["config"]["queues"]["Q2"]["servers"]
    ci = q["Q2"]["servers"]
    K2 = q["Q2"]["capacity"]
    rotas = ini["config"]["routes"]
    p_ret = rotas["Q3"][0]["probability"]
    sem = ini["sementes"]
    a1, b1 = q["Q1"]["service"]
    a2, b2 = q["Q2"]["service"]
    a3, b3 = q["Q3"]["service"]
    ea, eb = q["Q1"]["external_arrival"]
    ES = Ii["Q2"]["E_S"]
    mu_h = 60 * Ii["Q2"]["mu"]
    res_i, res_m = ini["resumo"]["Q2"], mel["resumo"]["Q2"]
    A, R, H = INTEGRANTES
    return [
        dict(titulo="Capa", orador=A, segundos=25, fala=(
            f"Olá, somos {A}, {R} e {H}. Neste trabalho modelamos uma assistência técnica "
            "hipotética como uma rede de três filas, simulamos com o simulador da T1 e testamos "
            "uma melhoria no reparo. Eu começo pelo problema, o Raul apresenta a análise do "
            "cenário inicial e o Humberto mostra a melhoria e a conclusão.")),
        dict(titulo="A realidade modelada", orador=A, segundos=40, fala=(
            "A realidade escolhida é uma assistência técnica de equipamentos. Cada cliente da rede "
            "é uma ordem de serviço, ou seja, um equipamento. Ele passa pela triagem, segue para o "
            "reparo e depois para o teste; se o teste reprova, volta ao reparo. É um cenário "
            "hipotético: os tempos são premissas nossas, não dados de uma empresa real. A pergunta "
            "é onde a rede trava e qual mudança simples reduz espera e perdas. Vamos ao desenho da rede.")),
        dict(titulo="Rede inicial", orador=A, segundos=60, fala=(
            f"A triagem, Q1, é G/G/1/{q['Q1']['capacity']}: um atendente, capacidade "
            f"{q['Q1']['capacity']} contando quem está em atendimento, serviço uniforme entre "
            f"{a1:g} e {b1:g} minutos. As chegadas externas ocorrem só na triagem, a cada {ea:g} a "
            f"{eb:g} minutos, com a primeira em {q['Q1']['first_arrival']:g}. Todo equipamento vai "
            f"para o reparo, Q2, G/G/{ci}/{K2}, com serviço de {a2:g} a {b2:g} minutos, e todo "
            f"reparado vai para o teste, Q3, G/G/1/{q['Q3']['capacity']}, de {a3:g} a {b3:g} minutos. "
            f"Do teste, {br(100 * p_ret, 0)}% retornam ao reparo, pela seta de retorno embaixo, e "
            f"{br(100 * (1 - p_ret), 0)}% saem da rede. Se a fila de destino está cheia, o "
            "equipamento é perdido, ou seja, encaminhado definitivamente para fora. A disciplina é FIFO.")),
        dict(titulo="Método de simulação", orador=A, segundos=50, fala=(
            f"Usamos o simulador da T1 com {br(ini['randoms'], 0)} números aleatórios por replicação. "
            f"Cada cenário teve {len(sem)} replicações, sementes {sem[0]} a {sem[-1]}, totalizando "
            f"{2 * len(sem)} execuções. A rede começa vazia e não descartamos aquecimento. Como cada "
            f"execução termina num tempo diferente, {br(ri['global_time'])} minutos no inicial e "
            f"{br(rm['global_time'])} no melhorado com a semente {sem[0]}, comparamos taxas por hora, "
            f"não contagens brutas. A semente {sem[0]} é o caso principal; o resumo das "
            f"{len(sem)} sementes vem no fim. Agora o Raul explica os cálculos.")),
        dict(titulo="Fórmulas e unidades (M9)", orador=R, segundos=55, fala=(
            "Obrigado, Artur. Usamos as fórmulas do M9. Pi i é a fração do tempo no estado i: tempo "
            "acumulado dividido pelo tempo total. A população média N é a soma de pi i vezes i, "
            "incluindo quem está em atendimento. O serviço médio é a mais b sobre dois e mi é o "
            "inverso; no estado i, a taxa mi i é o mínimo entre i e c, vezes mi. A vazão D é a soma "
            "de pi i vezes mi i, o mesmo que mi vezes B, o número médio de servidores ocupados. A "
            "utilização U é B sobre c, por servidor, e o tempo de resposta W é N sobre D, em minutos "
            "por visita. Conferimos as fórmulas com as imagens do M9 fornecidas pelo grupo. "
            f"No reparo, o serviço médio de {br(ES, 0)} minutos dá "
            f"{br(mu_h, 1)} atendimentos por hora por técnico; com {cm} técnicos, a capacidade "
            f"sobe para {br(cm * mu_h, 1)} por hora.")),
        dict(titulo="Probabilidades – cenário inicial", orador=R, segundos=45, fala=(
            "Esta tabela traz a probabilidade de todos os estados das três filas no cenário inicial; "
            "o traço indica estado acima da capacidade. Na triagem e no teste o sistema fica entre "
            f"zero e um equipamento. No reparo, a massa está no topo: {pct(Ii['Q2']['p'][K2 - 1], 4)} "
            f"do tempo com {K2 - 1} e {pct(Ii['Q2']['p'][K2], 4)} com {K2}, a capacidade máxima. "
            "Atenção: a fração de tempo com a fila cheia não é a probabilidade de perda, porque as "
            "chegadas não são Poisson.")),
        dict(titulo="Índices – cenário inicial", orador=R, segundos=45, fala=(
            f"No reparo, a utilização é {pct(Ii['Q2']['U'])} e a população média N é "
            f"{br(Ii['Q2']['L'])}. O tempo de resposta W é {br(Ii['Q2']['R'])} minutos por visita, contra "
            f"{br(Ii['Q1']['R'])} na triagem e {br(Ii['Q3']['R'])} no teste. Foram "
            f"{br(Ii['Q2']['perdas'], 0)} perdas no reparo em {br(ri['global_time'])} minutos, ou "
            f"{br(Ii['Q2']['perdas_h'], 4)} perdas por hora. Triagem e teste registraram "
            f"{br(Ii['Q1']['perdas'], 0)} e {br(Ii['Q3']['perdas'], 0)} perdas.")),
        dict(titulo="Diagnóstico: o reparo é o gargalo", orador=R, segundos=40, fala=(
            "O gargalo é o reparo. Ele recebe as ordens da triagem e ainda os retornos do teste, mas "
            f"um técnico atende em média {br(mu_h, 1)} por hora. Com utilização de "
            f"{pct(Ii['Q2']['U'])}, a fila fica cheia e as perdas aparecem. Triagem e teste têm folga, "
            f"com {pct(Ii['Q1']['U'])} e {pct(Ii['Q3']['U'])}. Por isso a melhoria deve atacar o "
            "reparo. Passo para o Humberto.")),
        dict(titulo="Rede melhorada", orador=H, segundos=30, fala=(
            f"Obrigado, Raul. Na rede melhorada mudamos um único parâmetro: o reparo passa de {ci} "
            f"para {cm} técnicos, virando G/G/{cm}/{K2}. A capacidade {K2} continua igual, assim como "
            f"chegadas, serviços e rotas, inclusive o retorno de {br(100 * p_ret, 0)}%.")),
        dict(titulo="Probabilidades – cenário melhorado", orador=H, segundos=25, fala=(
            f"Com {cm} técnicos, a distribuição do reparo se desloca para baixo: o estado vazio tem "
            f"{pct(Im['Q2']['p'][0], 4)} e o estado cheio aparece em apenas "
            f"{pct(Im['Q2']['p'][K2], 4)} do tempo. As regras da tabela são as mesmas do inicial.")),
        dict(titulo="Índices – cenário melhorado", orador=H, segundos=25, fala=(
            f"No melhorado, a utilização média por técnico do reparo é {pct(Im['Q2']['U'])}, a "
            f"população média N cai para {br(Im['Q2']['L'])} e o tempo de resposta W para "
            f"{br(Im['Q2']['R'])} minutos. Houve {br(Im['Q2']['perdas'], 0)} perdas em "
            f"{br(rm['global_time'])} minutos, {br(Im['Q2']['perdas_h'], 6)} por hora: raras, "
            "mas não eliminadas.")),
        dict(titulo="Comparação: reparo e efeito no teste", orador=H, segundos=40, fala=(
            f"No reparo, W varia {_variacao(Ii['Q2']['R'], Im['Q2']['R'])}, de {br(Ii['Q2']['R'])} "
            f"para {br(Im['Q2']['R'])} minutos por visita, e as perdas por hora variam "
            f"{_variacao(Ii['Q2']['perdas_h'], Im['Q2']['perdas_h'])}. A utilização por técnico cai "
            f"{br(100 * (Ii['Q2']['U'] - Im['Q2']['U']), 1)} pontos percentuais. O teste recebe mais "
            f"fluxo: a vazão estimada D de Q3 vai de {br(Ii['Q3']['X_h'])} para {br(Im['Q3']['X_h'])} "
            f"por hora e seu W de {br(Ii['Q3']['R'])} para {br(Im['Q3']['R'])} minutos. Não somamos "
            "os W das filas como tempo de rede, porque com retornos um equipamento pode visitar reparo "
            "e teste mais de uma vez; e o D do reparo inclui revisitas, não é o número de "
            "equipamentos únicos concluídos.")),
        dict(titulo="Robustez: 10 sementes", orador=H, segundos=25, fala=(
            f"Nas {len(sem)} sementes, o W do reparo ficou entre {br(res_i['R']['min'])} e "
            f"{br(res_i['R']['max'])} minutos no inicial, média {br(res_i['R']['media'])}, e entre "
            f"{br(res_m['R']['min'])} e {br(res_m['R']['max'])} no melhorado, média "
            f"{br(res_m['R']['media'])}. As perdas por hora seguem o mesmo sentido. É um resumo "
            "descritivo, sem intervalo de confiança.")),
        dict(titulo="Conclusão, custo e limites", orador=H, segundos=30, fala=(
            "Neste modelo hipotético, o segundo técnico no reparo reduz fortemente espera e perdas. "
            "O custo é um técnico adicional, que não estimamos em dinheiro. A simulação não prova "
            "rentabilidade nem que uma empresa real se comportaria assim. Limites: modelo hipotético, "
            "tempos uniformes supostos, sem aquecimento, e a vazão D é estimada pela ocupação, não "
            "contada.")),
        dict(titulo="Referências e próximos passos", orador=H, segundos=15, fala=(
            "As referências são o M9 do professor Afonso Sales, o simulador da T1, a análise da T2 e "
            "o texto do MIT sobre a lei de Little. Obrigado pela atenção.")),
    ]


# ---------------------------------------------------------------- PPTX

def _texto(slide, x, y, w, h, linhas, size=20, cor=TEXTO, bold=False, align=PP_ALIGN.LEFT):
    caixa = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caixa.text_frame
    tf.word_wrap = True
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linha
        p.alignment = align
        p.space_after = Pt(6)
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, cor
    return caixa


def _forma(slide, x, y, w, h, linhas, fundo, cor=BRANCO, size=15):
    f = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    f.fill.solid()
    f.fill.fore_color.rgb = fundo
    f.line.color.rgb = NAVY
    tf = f.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linha
        p.alignment = PP_ALIGN.CENTER
        p.font.size, p.font.color.rgb, p.font.bold = Pt(size + 2 if i == 0 else size), cor, i == 0
    return f


def _seta(slide, x1, y1, x2, y2, cor=NAVY, ponta=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = cor
    c.line.width = Pt(2.5)
    if ponta:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def _tabela(slide, x, y, w, h, linhas, size=14, destaque=()):
    tab = slide.shapes.add_table(len(linhas), len(linhas[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for i, linha in enumerate(linhas):
        for j, valor in enumerate(linha):
            cel = tab.cell(i, j)
            cel.text = str(valor)
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            cel.fill.solid()
            cel.fill.fore_color.rgb = NAVY if i == 0 else AMARELO if i in destaque else LISTRA if i % 2 else BRANCO
            p = cel.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            p.font.size, p.font.bold = Pt(size), i == 0
            p.font.color.rgb = BRANCO if i == 0 else TEXTO
    return tab


def _barras(slide, x, y, w, h, titulo, categorias, series, formato="0.0"):
    dados = CategoryChartData()
    dados.categories = categorias
    for nome, valores in series.items():
        dados.add_series(nome, valores)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), dados
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = titulo
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(16)
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format, plot.data_labels.number_format_is_linked = formato, False
    plot.data_labels.font.size = Pt(12)
    for serie, cor in zip(plot.series, (NAVY, AMARELO)):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = cor
    return chart


def _slide(prs, n, total, titulo, fala):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = CLARO
    _texto(s, 0.6, 0.3, 12.1, 0.9, [titulo], 32, NAVY, True)
    barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.15), Inches(1.6), Inches(0.07))
    barra.fill.solid()
    barra.fill.fore_color.rgb = AMARELO
    barra.line.fill.background()
    _texto(s, 11.3, 7.0, 1.6, 0.4, [f"{n}/{total}"], 12, TEAL, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = fala
    return s


def _diagrama(s, config, y0=1.6, destaque_q2=False):
    q, rotas = config["queues"], config["routes"]
    xs = {"Q1": 1.7, "Q2": 5.4, "Q3": 9.1}
    w, h = 2.4, 1.5
    meio = y0 + h / 2
    for nome, x in xs.items():
        f = q[nome]
        c, (a, b) = f["servers"], f["service"]
        amarelo = destaque_q2 and nome == "Q2"
        _forma(s, x, y0, w, h, [
            f"{nome} – {PAPEIS[nome]}", f"G/G/{c}/{f['capacity']}",
            f"{c} servidor{'es' if c > 1 else ''}", f"serviço U({a:g}, {b:g}) min",
        ], AMARELO if amarelo else NAVY, NAVY if amarelo else BRANCO, 13)
    ea, eb = q["Q1"]["external_arrival"]
    _texto(s, 0.1, y0 - 0.35, 1.6, 1.0,
           ["chegada externa", f"U({ea:g}, {eb:g}) min", f"1ª em {q['Q1']['first_arrival']:g} min"], 11, TEAL, True,
           PP_ALIGN.CENTER)
    _seta(s, 0.25, meio, xs["Q1"], meio)
    for origem, destino in (("Q1", "Q2"), ("Q2", "Q3")):
        p = next(r["probability"] for r in rotas[origem] if r["to"] == destino)
        x1, x2 = xs[origem] + w, xs[destino]
        _texto(s, x1, meio - 0.5, x2 - x1, 0.4, [f"{br(100 * p, 0)}%"], 14, NAVY, True, PP_ALIGN.CENTER)
        _seta(s, x1, meio, x2, meio)
    p_ret = next(r["probability"] for r in rotas["Q3"] if r["to"] == "Q2")
    x_saida = xs["Q3"] + w
    _texto(s, x_saida, meio - 0.55, 1.6, 0.45, [f"{br(100 * (1 - p_ret), 0)}% saída"], 14, NAVY, True,
           PP_ALIGN.CENTER)
    _seta(s, x_saida, meio, 12.9, meio)
    # Retorno Q3 -> Q2 por baixo das caixas: desce, vai à esquerda, sobe com ponta em Q2.
    baixo, y_ret = y0 + h, y0 + h + 0.8
    x3, x2 = xs["Q3"] + w / 2, xs["Q2"] + w / 2
    _seta(s, x3, baixo, x3, y_ret, TEAL, False)
    _seta(s, x3, y_ret, x2, y_ret, TEAL, False)
    _seta(s, x2, y_ret, x2, baixo, TEAL)
    _texto(s, x2 + 0.2, y_ret + 0.05, x3 - x2 - 0.4, 0.45,
           [f"{br(100 * p_ret, 0)}% retorno (retrabalho) Q3 → Q2"], 14, TEAL, True, PP_ALIGN.CENTER)
    _texto(s, 0.6, y_ret + 0.75, 12.1, 0.9, [
        "Legenda: caixa = fila G/G/c/K (c servidores; K inclui quem está em serviço) · seta = rota e "
        "probabilidade · tempos em minutos · FIFO · Q2 e Q3 sem chegadas externas · chegada a fila "
        "cheia = perda (encaminhamento externo definitivo)."], 13, TEXTO)


def _tabela_probs(s, cenario, rep):
    K = max(f["capacity"] for f in rep["queues"].values())
    linhas = [["n"] + [f"{f} (K={rep['queues'][f]['capacity']})" for f in FILAS]]
    for n in range(K + 1):
        linhas.append([str(n)] + [pct(rep["indices"][f]["p"][n], 4) if n < len(rep["indices"][f]["p"]) else "–"
                                  for f in FILAS])
    linhas.append(["Soma"] + [pct(sum(rep["indices"][f]["p"]), 4) for f in FILAS])
    _tabela(s, 0.6, 1.45, 8.4, 5.4, linhas, 13)
    _texto(s, 9.3, 1.45, 3.6, 5.4, [
        f"Cenário {cenario}, semente {rep['seed']}.",
        "p_n = tempo no estado n / T.",
        "“–” = estado acima da capacidade K.",
        "Percentuais com 4 casas; a soma arredondada pode divergir de 100%.",
        "P(K) é fração do tempo com a fila cheia, não a probabilidade de perda (chegadas não Poisson).",
    ], 15)


def _tabela_indices(s, cenario, rep, destaque=()):
    linhas = [["Fila", "c", "N (clientes)", "U", "D_h (atend./h)", "W (min/visita)", "Perdas", "Perdas/h"]]
    for f in FILAS:
        i = rep["indices"][f]
        linhas.append([f"{f} – {PAPEIS[f]}", str(rep["queues"][f]["servers"]), br(i["L"], 4), pct(i["U"]),
                       br(i["X_h"], 4), br(i["R"], 4) if i["R"] is not None else "–", br(i["perdas"], 0),
                       br(i["perdas_h"], 6)])
    _tabela(s, 0.6, 1.5, 12.1, 2.4, linhas, 15, destaque)
    _texto(s, 0.6, 4.2, 12.1, 2.6, [
        f"Cenário {cenario}, semente {rep['seed']}: T = {br(rep['global_time'])} min simulados.",
        "U = B/c (média por servidor); D = μ·B (atend./min), D_h = 60·D: estimativa da vazão local, inclui revisitas.",
        "W = N/D, com D por minuto (Little): estimativa de longo prazo por visita, não por ordem.",
        "Perdas/h = 60 · perdas / T: taxa observada, comparável entre execuções de duração diferente.",
    ], 17)


def apresentacao(dados, falas, caminho):
    ini, mel = dados["inicial"], dados["melhorado"]
    ri, rm = principal(ini), principal(mel)
    Ii, Im = ri["indices"], rm["indices"]
    q2i, q2m = ini["config"]["queues"]["Q2"], mel["config"]["queues"]["Q2"]
    sem = ini["sementes"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    total = len(falas)
    it = iter(enumerate(falas, 1))

    def proximo():
        n, f = next(it)
        return _slide(prs, n, total, f["titulo"], f["fala"])

    # 1 capa (navy)
    s = proximo()
    s.background.fill.fore_color.rgb = NAVY
    for forma in list(s.shapes)[:2]:
        forma._element.getparent().remove(forma._element)
    _texto(s, 0.9, 1.6, 11.5, 1.4, ["Rede de filas de uma assistência técnica"], 44, BRANCO, True)
    _texto(s, 0.9, 3.0, 11.5, 0.8, ["T2 – Métodos Analíticos · simulação com o simulador da T1"], 24, AMARELO)
    _texto(s, 0.9, 4.3, 11.5, 1.8, list(INTEGRANTES), 24, BRANCO)
    _texto(s, 0.9, 6.3, 11.5, 0.6, ["Cenário hipotético: triagem → reparo → teste, com retrabalho"], 16, BRANCO)

    # 2 realidade
    s = proximo()
    _texto(s, 0.6, 1.5, 12.1, 5.3, [
        "• Assistência técnica hipotética de equipamentos; cliente = ordem de serviço / equipamento.",
        "• Fluxo: triagem → reparo → teste; reprovado no teste volta ao reparo (retrabalho).",
        "• Fila cheia no destino = equipamento perdido (encaminhado definitivamente para fora).",
        "• Tempos são premissas do grupo, não dados de uma empresa real.",
        "• Pergunta: onde a rede trava e que mudança simples reduz espera e perdas?",
    ], 22)

    # 3 rede inicial
    s = proximo()
    _diagrama(s, ini["config"])

    # 4 método
    s = proximo()
    _texto(s, 0.6, 1.5, 12.1, 5.3, [
        f"• Simulador da T1, {br(ini['randoms'], 0)} aleatórios por replicação (critério de parada).",
        f"• {len(sem)} replicações por cenário, sementes {sem[0]}..{sem[-1]}: {2 * len(sem)} execuções.",
        "• Rede inicia vazia; sem aquecimento (warmup) descartado.",
        f"• Semente {sem[0]}: T = {br(ri['global_time'])} min (inicial) e {br(rm['global_time'])} min (melhorado).",
        "• Tempos finais diferentes → comparar taxas por hora, não contagens brutas.",
        f"• Slides usam a semente {sem[0]}; resumo das {len(sem)} sementes no slide de robustez.",
    ], 22)

    # 5 fórmulas
    s = proximo()
    _tabela(s, 0.6, 1.45, 8.3, 5.5, [
        ["Índice", "Fórmula (M9, i = 1..K)", "Unidade"],
        ["π_i", "tempo no estado i / T", "fração do tempo"],
        ["N", "Σ π_i · i", "clientes"],
        ["μ", "1 / E[S]; E[S] = (a + b)/2", "atend./min/servidor"],
        ["μ_i", "min(i, c) · μ", "atend./min"],
        ["D", "Σ π_i · μ_i = μ · B", "atend./min"],
        ["D_h", "60 · D", "atend./h"],
        ["U", "Σ π_i · min(i, c)/c = B/c", "por servidor"],
        ["W", "N / D (lei de Little)", "min/visita"],
        ["B (auxiliar)", "Σ π_i · min(i, c)", "servidores ocupados"],
        ["Perdas/h", "60 · perdas / T", "perdas/h"],
    ], 13)
    _forma(s, 9.2, 1.45, 3.7, 3.4, [
        "Exemplo: reparo (Q2)",
        f"E[S] = ({q2i['service'][0]:g} + {q2i['service'][1]:g}) / 2 = {br(Ii['Q2']['E_S'], 0)} min",
        f"μ = {br(Ii['Q2']['mu'], 3)}/min = {br(60 * Ii['Q2']['mu'], 1)}/h por técnico",
        f"{q2m['servers']} técnicos: capacidade {br(60 * q2m['servers'] * Im['Q2']['mu'], 1)}/h",
    ], TEAL, BRANCO, 14)
    _texto(s, 9.2, 5.0, 3.7, 1.9, [
        "Sem M/M/c analítico: serviços e chegadas são uniformes.",
        "Fórmulas conferidas com as imagens do M9 fornecidas pelo grupo.",
    ], 13)

    # 6 probabilidades inicial
    _tabela_probs(proximo(), "inicial", ri)

    # 7 índices inicial
    _tabela_indices(proximo(), "inicial", ri, destaque=(2,))

    # 8 diagnóstico
    s = proximo()
    _texto(s, 0.6, 1.5, 6.6, 5.3, [
        f"• Q2 recebe triagem + retornos do teste; 1 técnico atende ~{br(60 * Ii['Q2']['mu'], 1)}/h.",
        f"• U(Q2) = {pct(Ii['Q2']['U'])}; fila cheia {pct(Ii['Q2']['p'][q2i['capacity']])} do tempo.",
        f"• {br(Ii['Q2']['perdas'], 0)} perdas em Q2 → {br(Ii['Q2']['perdas_h'], 4)} perdas/h.",
        f"• Folga em Q1 ({pct(Ii['Q1']['U'])}) e Q3 ({pct(Ii['Q3']['U'])}).",
        "• Conclusão: a melhoria deve atacar o reparo.",
    ], 19)
    _barras(s, 7.5, 1.5, 5.3, 5.0, "Utilização por fila – inicial (%)", list(FILAS),
            {"Inicial": [100 * Ii[f]["U"] for f in FILAS]})

    # 9 rede melhorada
    s = proximo()
    _diagrama(s, mel["config"], destaque_q2=True)
    _forma(s, 2.4, 6.45, 8.5, 0.75, [
        f"Única mudança: Q2 de {q2i['servers']} para {q2m['servers']} técnicos; capacidade "
        f"{q2m['capacity']} mantida."], AMARELO, NAVY, 15)

    # 10 probabilidades melhorado
    _tabela_probs(proximo(), "melhorado", rm)

    # 11 índices melhorado
    _tabela_indices(proximo(), "melhorado", rm, destaque=(2,))

    # 12 comparação
    s = proximo()
    linhas = [["Métrica", "Inicial", "Melhorado", "Variação"]]
    for f, chave, rotulo, d in (("Q2", "L", "Q2 N (clientes)", 4), ("Q2", "R", "Q2 W (min/visita)", 4),
                                ("Q2", "X_h", "Q2 D_h (atend./h)", 4), ("Q2", "perdas_h", "Q2 perdas/h", 6),
                                ("Q3", "X_h", "Q3 D_h (atend./h)", 4), ("Q3", "R", "Q3 W (min/visita)", 4)):
        i, m = Ii[f][chave], Im[f][chave]
        linhas.append([rotulo, br(i, d), br(m, d), _variacao(i, m)])
    linhas.insert(2, ["Q2 U (por técnico)", pct(Ii["Q2"]["U"]), pct(Im["Q2"]["U"]),
                      br(100 * (Im["Q2"]["U"] - Ii["Q2"]["U"]), 2) + " p.p."])
    _tabela(s, 0.6, 1.45, 7.0, 4.0, linhas, 14)
    _barras(s, 7.9, 1.45, 4.9, 4.0, "W por visita (min)", list(FILAS),
            {"Inicial": [Ii[f]["R"] or 0 for f in FILAS], "Melhorado": [Im[f]["R"] or 0 for f in FILAS]})
    _texto(s, 0.6, 5.65, 12.1, 1.3, [
        "Não somar W das filas como tempo na rede: com retornos, uma ordem visita Q2/Q3 mais de uma vez.",
        "D de Q2 inclui revisitas; não é o número de equipamentos únicos concluídos.",
    ], 15)

    # 13 robustez
    s = proximo()
    linhas = [["Cenário", "W Q2 média", "W Q2 mín", "W Q2 máx", "Perdas/h média", "Perdas/h mín", "Perdas/h máx"]]
    for c in CENARIOS:
        r = dados[c]["resumo"]["Q2"]
        linhas.append([c] + [br(r["R"][k], 4) for k in ("media", "min", "max")]
                      + [br(r["perdas_h"][k], 6) for k in ("media", "min", "max")])
    _tabela(s, 0.6, 1.5, 12.1, 1.8, linhas, 15)
    _texto(s, 0.6, 3.7, 12.1, 3.0, [
        f"• {len(sem)} sementes por cenário ({sem[0]}..{sem[-1]}); W em min/visita, perdas por hora.",
        "• O sentido da melhoria se repete em todas as sementes do resumo acima.",
        "• Resumo descritivo (média, mínimo, máximo); não é intervalo de confiança.",
    ], 19)

    # 14 conclusão
    s = proximo()
    _texto(s, 0.6, 1.5, 12.1, 5.3, [
        f"• Com {q2m['servers']} técnicos no reparo, W de Q2: {br(Ii['Q2']['R'])} → {br(Im['Q2']['R'])} min/visita.",
        f"• Perdas/h em Q2: {br(Ii['Q2']['perdas_h'], 4)} → {br(Im['Q2']['perdas_h'], 6)} (raras, não eliminadas).",
        "• Custo: um técnico adicional (não estimado em dinheiro).",
        "• A simulação não prova rentabilidade nem o comportamento de uma empresa real.",
        "• Limites: modelo hipotético, tempos uniformes supostos, sem warmup, D estimada pela ocupação.",
    ], 21)

    # 15 referências
    s = proximo()
    _texto(s, 0.6, 1.5, 12.1, 5.3, [
        "• Prof. Afonso Sales, M9 – Avaliação de um Modelo Simulado, PUCRS 2026/2 (imagens fornecidas pelo grupo).",
        "• T1/simulador_fila.py – motor de simulação reutilizado.",
        "• T2/analise.py – índices (π_i, N, U, D, W, perdas/h) e resumo das sementes.",
        f"• Larson & Odoni, Urban Operations Research, seção 4.4 (lei de Little): {MIT}",
        "• Próximos passos: gravar o vídeo, publicar no YouTube (não listado), registrar o link "
        "em VIDEO.txt e reconstruir o ZIP.",
    ], 19)
    prs.save(caminho)


# ---------------------------------------------------------------- XLSX

def planilha(dados, caminho):
    wb = Workbook()
    negrito, cabec = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")

    def cabecalho(ws, linha, nomes):
        for j, nome in enumerate(nomes, 1):
            cel = ws.cell(linha, j, nome)
            cel.font, cel.fill, cel.alignment = negrito, cabec, Alignment(wrap_text=True, vertical="center")
        for j in range(1, len(nomes) + 1):
            ws.column_dimensions[ws.cell(linha, j).column_letter].width = 17

    leia = wb.active
    leia.title = "Leia-me"
    leia.column_dimensions["A"].width = 110
    s0 = dados["inicial"]["semente_principal"]
    for i, linha in enumerate([
        "T2 – Rede de filas de uma assistência técnica (modelo hipotético)",
        "Integrantes: " + ", ".join(INTEGRANTES),
        f"Fonte: T2/resultados/inicial.json e melhorado.json; abas Estados_* e Indices usam a semente {s0}.",
        "Premissas: cliente = ordem de serviço; FIFO; K inclui quem está em serviço; chegada a fila cheia = perda.",
        "Unidades: tempos em minutos; D em atendimentos/min e D_h = 60*D em atendimentos/h; W em min/visita; perdas_h em perdas/h.",
        "Notação do M9 (Prof. Afonso Sales, PUCRS 2026/2), i = 1..K: pi_i = tempo_i / T (coluna p_n); N = SUM(pi_i*i); "
        "mu = 1/E[S], E[S] = (a+b)/2; mu_i = MIN(i,c)*mu; D = SUM(pi_i*mu_i) = mu*B; U = SUM(pi_i*MIN(i,c)/c) = B/c; W = N/D.",
        "B = SUM(MIN(n,c)*p_n) é auxiliar (servidores ocupados); o estado 0 contribui 0. perdas_h = 60*perdas/T.",
        "Dicionário com o JSON/CSV (chaves inalteradas): N = L, D = X, D_h = X_h, W = R; U, B, perdas, perdas_h iguais.",
        "Fórmulas conferidas com as imagens do M9 fornecidas pelo grupo.",
        "D é estimativa da vazão local pela ocupação (inclui revisitas Q3→Q2), não contagem; W é estimativa de longo prazo por visita.",
        "P(K) não é a probabilidade de perda (chegadas não Poisson).",
        "Replicacoes: valores exportados do JSON; resumo (média/mín/máx) por fórmulas, descritivo, sem IC.",
    ], 1):
        leia.cell(i, 1, linha).font = Font(bold=i == 1, size=14 if i == 1 else 11)

    linhas_ind = {}
    ind = wb.create_sheet("Indices")
    cols = ["Cenário", "Fila", "c", "T (min)", "a (min)", "b (min)", "Perdas", "Soma p_n", "N", "B", "U",
            "E[S] (min)", "mu (/min/servidor)", "D (/min)", "D_h (/h)", "W (min/visita)", "Perdas_h (/h)"]
    cabecalho(ind, 1, cols)
    ind.freeze_panes = "C2"
    r_ind = 2
    for cen in CENARIOS:
        rep = principal(dados[cen])
        ws = wb.create_sheet(f"Estados_{cen}", 1 + CENARIOS.index(cen))
        ws["A1"], ws["B1"] = "Cenário", cen
        ws["A2"], ws["B2"] = "T (min)", rep["global_time"]
        ws["A3"], ws["B3"] = "Semente", rep["seed"]
        cabecalho(ws, 5, ["Fila", "c", "a (min)", "b (min)", "Perdas", "n", "Tempo_n (min)", "p_n", "n*p_n",
                          "min(n,c)*p_n"])
        ws.freeze_panes = "A6"
        r = 6
        for fila in FILAS:
            dados_fila = rep["queues"][fila]
            a, b = dados[cen]["config"]["queues"][fila]["service"]
            r1 = r
            for n, tempo in enumerate(dados_fila["accumulated"]):
                for j, v in enumerate([fila, dados_fila["servers"], a, b, dados_fila["losses"], n, tempo], 1):
                    ws.cell(r, j, v)
                ws.cell(r, 8, f"=G{r}/$B$2").number_format = "0.0000%"
                ws.cell(r, 9, f"=F{r}*H{r}").number_format = "0.000000"
                ws.cell(r, 10, f"=MIN(F{r},B{r})*H{r}").number_format = "0.000000"
                r += 1
            e = f"'Estados_{cen}'!"
            r2 = r - 1
            formulas = [cen, fila, f"={e}B{r1}", f"={e}$B$2", f"={e}C{r1}", f"={e}D{r1}", f"={e}E{r1}",
                        f"=SUM({e}H{r1}:H{r2})", f"=SUM({e}I{r1}:I{r2})", f"=SUM({e}J{r1}:J{r2})",
                        f"=J{r_ind}/C{r_ind}", f"=(E{r_ind}+F{r_ind})/2", f"=1/L{r_ind}", f"=M{r_ind}*J{r_ind}",
                        f"=60*N{r_ind}", f'=IF(N{r_ind}>0,I{r_ind}/N{r_ind},"")', f"=60*G{r_ind}/D{r_ind}"]
            for j, v in enumerate(formulas, 1):
                ind.cell(r_ind, j, v).number_format = "0.000000"
            ind.cell(r_ind, 11).number_format = "0.00%"
            linhas_ind[(cen, fila)] = r_ind
            r_ind += 1

    comp = wb.create_sheet("Comparacao")
    cabecalho(comp, 1, ["Fila", "Métrica", "Inicial", "Melhorado", "Variação", "Unidade da variação"])
    comp.freeze_panes = "A2"
    r = 2
    for fila in FILAS:
        for rotulo, col in (("N", "I"), ("U", "K"), ("D_h (/h)", "O"), ("W (min/visita)", "P"),
                            ("Perdas", "G"), ("Perdas_h (/h)", "Q")):
            ri, rm = linhas_ind[("inicial", fila)], linhas_ind[("melhorado", fila)]
            comp.cell(r, 1, fila)
            comp.cell(r, 2, rotulo)
            comp.cell(r, 3, f"=Indices!{col}{ri}").number_format = "0.000000"
            comp.cell(r, 4, f"=Indices!{col}{rm}").number_format = "0.000000"
            if col == "K":
                comp.cell(r, 5, f"=(D{r}-C{r})*100").number_format = "0.00"
                comp.cell(r, 6, "pontos percentuais")
            else:
                comp.cell(r, 5, f'=IF(C{r}=0,"",(D{r}-C{r})/C{r})').number_format = "0.00%"
                comp.cell(r, 6, "% relativo ao inicial")
            r += 1

    rep_ws = wb.create_sheet("Replicacoes")
    metricas = ["L", "B", "U", "X_h", "R", "perdas", "perdas_h"]
    rotulos = {"L": "N", "X_h": "D_h", "R": "W"}  # notação M9; chaves JSON inalteradas
    cabecalho(rep_ws, 1, ["Cenário", "Semente", "Fila", "T (min)"] + [rotulos.get(m, m) for m in metricas])
    rep_ws.freeze_panes = "A2"
    r, blocos = 2, []
    for cen in CENARIOS:
        for fila in FILAS:
            r1 = r
            for rep in dados[cen]["replicacoes"]:
                i = rep["indices"][fila]
                for j, v in enumerate([cen, rep["seed"], fila, rep["global_time"]]
                                      + [i[m] if i[m] is not None else "" for m in metricas], 1):
                    rep_ws.cell(r, j, v)
                r += 1
            blocos.append((cen, fila, r1, r - 1))
    r += 1
    rep_ws.cell(r, 1, "Resumo por fórmulas (descritivo, sem intervalo de confiança)").font = Font(bold=True)
    r += 1
    cabecalho(rep_ws, r, ["Cenário", "Fila", "Métrica", "Média", "Mín", "Máx", "n"])
    r += 1
    for cen, fila, r1, r2 in blocos:
        for k, m in enumerate(metricas):
            col = rep_ws.cell(1, 5 + k).column_letter
            faixa = f"{col}{r1}:{col}{r2}"
            for j, v in enumerate([cen, fila, rotulos.get(m, m), f"=AVERAGE({faixa})", f"=MIN({faixa})", f"=MAX({faixa})",
                                   f"=COUNT({faixa})"], 1):
                rep_ws.cell(r, j, v).number_format = "0.000000"
            r += 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(caminho)


# ---------------------------------------------------------------- roteiro MD/PDF

def _md(falas):
    linhas = ["# Roteiro do vídeo – T2", "",
              f"Integrantes: {', '.join(INTEGRANTES)}. Duração-alvo: "
              f"{sum(f['segundos'] for f in falas) // 60} min {sum(f['segundos'] for f in falas) % 60} s.", ""]
    for n, f in enumerate(falas, 1):
        linhas += [f"## Slide {n} – {f['titulo']}", "", f"**{f['orador']}** (~{f['segundos']} s)", "", f["fala"], ""]
    linhas += ["## Checklist de gravação", ""] + [f"- [ ] {c}" for c in CHECKLIST]
    linhas += ["", "O vídeo ainda não foi gravado nem publicado; o link real entra em T2/entrega/VIDEO.txt."]
    return "\n".join(linhas) + "\n"


def _fontes():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/arialbd.ttf"))
        return "Arial", "Arial-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"  # ponytail: sem Arial perde acentos fora do Latin-1


def _pdf(falas, caminho):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    normal, forte = _fontes()
    corpo = ParagraphStyle("corpo", fontName=normal, fontSize=11, leading=15)
    titulo = ParagraphStyle("titulo", fontName=forte, fontSize=18, leading=22, spaceAfter=8)
    sub = ParagraphStyle("sub", fontName=forte, fontSize=13, leading=17, spaceBefore=10, spaceAfter=4)
    total = sum(f["segundos"] for f in falas)
    hist = [Paragraph("Roteiro do vídeo – T2", titulo),
            Paragraph(escape(f"Integrantes: {', '.join(INTEGRANTES)}. Duração-alvo: {total // 60} min "
                             f"{total % 60} s."), corpo)]
    for n, f in enumerate(falas, 1):
        hist += [Paragraph(escape(f"Slide {n} – {f['titulo']}"), sub),
                 Paragraph(escape(f"{f['orador']} (~{f['segundos']} s)"), corpo), Spacer(1, 3),
                 Paragraph(escape(f["fala"]), corpo)]
    hist.append(Paragraph("Checklist de gravação", sub))
    hist += [Paragraph(escape(f"[ ] {c}"), corpo) for c in CHECKLIST]
    hist += [Spacer(1, 8), Paragraph("O vídeo ainda não foi gravado nem publicado; o link real entra em "
                                     "T2/entrega/VIDEO.txt.", corpo)]
    SimpleDocTemplate(str(caminho), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                      bottomMargin=2 * cm, title="Roteiro T2").build(hist)


# ---------------------------------------------------------------- orquestração

def gerar(destino=ENTREGA, resultados=RESULTADOS):
    """Gera os artefatos em destino; devolve {nome: caminho}."""
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    dados = carregar(resultados)
    falas = roteiro(dados)
    saidas = {n: destino / n for n in ("Apresentacao_T2.pptx", "Indices_T2.xlsx", "Roteiro_T2.pdf",
                                       "Roteiro_T2.md", "VIDEO_PENDENTE.txt")}
    apresentacao(dados, falas, saidas["Apresentacao_T2.pptx"])
    planilha(dados, saidas["Indices_T2.xlsx"])
    _pdf(falas, saidas["Roteiro_T2.pdf"])
    saidas["Roteiro_T2.md"].write_text(_md(falas), encoding="utf-8")
    saidas["VIDEO_PENDENTE.txt"].write_text(
        "Vídeo ainda não gravado nem publicado.\n"
        "Após gravar e publicar no YouTube (não listado), crie T2/entrega/VIDEO.txt com o link real,\n"
        "apague este arquivo e reconstrua o ZIP: python T2/gerar_entrega.py zip\n", encoding="utf-8")
    return saidas


def empacotar(raiz=BASE.parent):
    """Cria raiz/T2/Entrega_T2.zip com T2/ (sem lixo) e T1/simulador_fila.py; exige o PDF da apresentação."""
    raiz = Path(raiz)
    t2 = raiz / "T2"
    pdf = t2 / "entrega" / "Apresentacao_T2.pdf"
    if not pdf.is_file():
        raise FileNotFoundError(f"exporte a apresentação para PDF antes de empacotar: {pdf}")
    destino = t2 / "Entrega_T2.zip"
    arquivos = [p for p in sorted(t2.rglob("*")) if p.is_file() and p != destino and p.suffix != ".pyc"
                and not p.name.startswith("~$") and not IGNORAR & set(p.relative_to(t2).parts)]
    arquivos.append(raiz / "T1" / "simulador_fila.py")
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        for p in arquivos:
            z.write(p, p.relative_to(raiz).as_posix())
    return destino


if __name__ == "__main__":
    if sys.argv[1:] == ["zip"]:
        print(empacotar())
    else:
        for caminho in gerar().values():
            print(caminho)
